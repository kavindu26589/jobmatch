"""Upload cache and format-preserving CV export.

Uploaded resumes are stored on disk (``<REPORT_DIR>/uploads/<id>/``) so the UI
can regenerate the edited CV in the *same* file format it was uploaded in
(``.pdf``, ``.docx``, ``.txt``, ``.md``, ``.rtf``).

Editing is normally done on the extracted text, but PDF/DOCX exports keep the
original design: the changed phrases are located and rewritten *in place* on
the original file (see :mod:`jobmatcher.pdfpatch`). A clean reflow regenerate
is used as a fallback when the changes cannot be positioned, and for the plain
text formats.
"""

from __future__ import annotations

import io
import re
import shutil
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, TypeAlias, cast

MAX_UPLOAD_BYTES = 5 * 1024 * 1024

_BULLET_RE = re.compile(r"^\s*([-*•])\s+")
_HEADER_RE = re.compile(r"^[A-Z][A-Za-z &/+.-]{2,48}:?$")

CONTENT_TYPES: dict[str, str] = {
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/markdown; charset=utf-8",
    ".rtf": "application/rtf",
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ".pdf": "application/pdf",
}


@dataclass(frozen=True)
class UploadRecord:
    upload_id: str
    filename: str
    extension: str
    path: Path


def _clean_filename(name: str) -> str:
    return Path(name or "resume").name


class UploadStore:
    """Persists uploaded files keyed by a random id."""

    def __init__(self, base_dir: str | Path) -> None:
        self.base = Path(base_dir) / "uploads"

    def _record_dir(self, upload_id: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{32}", upload_id):
            raise KeyError(upload_id)
        return self.base / upload_id

    def save(self, filename: str, data: bytes) -> UploadRecord:
        """Store raw bytes under a fresh random id and return its record."""
        filename = _clean_filename(filename)
        extension = Path(filename).suffix.lower()
        upload_id = uuid.uuid4().hex
        record_dir = self._record_dir(upload_id)
        record_dir.mkdir(parents=True, exist_ok=True)
        (record_dir / f"original{extension}").write_bytes(data)
        (record_dir / "meta.txt").write_text(f"{filename}\n{extension}\n", encoding="utf-8")
        return UploadRecord(
            upload_id=upload_id,
            filename=filename,
            extension=extension,
            path=record_dir / f"original{extension}",
        )

    def record(self, upload_id: str) -> UploadRecord:
        """Return the stored record, raising ``KeyError`` if unknown."""
        record_dir = self._record_dir(upload_id)
        meta = record_dir / "meta.txt"
        if not meta.exists():
            raise KeyError(upload_id)
        filename, extension = meta.read_text(encoding="utf-8").splitlines()[:2]
        return UploadRecord(
            upload_id=upload_id,
            filename=filename,
            extension=extension,
            path=record_dir / f"original{extension}",
        )

    def render(
        self, upload_id: str, text: str, as_format: str | None = None
    ) -> tuple[bytes, str, str, list[str]]:
        """Render edited ``text``, preserving the original layout when possible.

        For PDF/DOCX uploads exported in their *original* format, changes are
        patched in place on the original file (design preserved) and a fallback
        reflow is produced only when the patch cannot place the changes. Other
        cases use the clean reflow renderers.

        Returns ``(bytes, filename, content_type, notes)`` where ``notes``
        describe anything the in-place patch could not do.
        """
        record = self.record(upload_id)
        extension = (as_format or record.extension).strip().lower()
        if not extension.startswith("."):
            extension = "." + extension
        if extension not in CONTENT_TYPES:
            raise ValueError(
                f"unsupported export format '{extension}' "
                f"(supported: {', '.join(sorted(CONTENT_TYPES))})"
            )
        name = Path(record.filename).stem or "resume"
        notes: list[str] = []
        if (
            extension == record.extension
            and extension in (".pdf", ".docx")
            and record.path.exists()
        ):
            try:
                from ..pdfpatch import patch_binary
                from ..resume.parsers import extract_text

                original = extract_text(record.path)
                result = patch_binary(
                    record.path.read_bytes(),
                    original,
                    text,
                    cast(Literal["pdf", "docx"], extension.lstrip(".")),
                )
                notes.extend(result.notes)
                if result.total == 0:
                    return result.data, f"{name}{extension}", CONTENT_TYPES[extension], notes
                if result.applied:
                    notes.insert(0, f"patched in place ({result.applied}/{result.total} changes)")
                    return result.data, f"{name}{extension}", CONTENT_TYPES[extension], notes
                notes.insert(0, "reflowed: could not position the changes on the original file")
            except Exception as exc:  # pragma: no cover - defensive fallback
                notes.append(f"in-place patch failed ({type(exc).__name__}); reflowed")
        module = _EXTENSION_RENDERERS[extension]
        data = module(text, name)
        return data, f"{name}{extension}", CONTENT_TYPES[extension], notes

    def delete(self, upload_id: str) -> None:
        record_dir = self._record_dir(upload_id)
        if record_dir.exists():
            shutil.rmtree(record_dir, ignore_errors=True)


# ---------------------------------------------------------------------------
# Renderers
# ---------------------------------------------------------------------------


def _to_runs(text: str) -> list[tuple[str, bool]]:
    """Split text into (line, looks_like_heading) runs, skipping blanks."""
    runs: list[tuple[str, bool]] = []
    for raw in text.strip().splitlines():
        line = raw.strip()
        if not line:
            continue
        runs.append((line, bool(_HEADER_RE.fullmatch(line))))
    return runs


def render_txt(text: str, name: str) -> bytes:
    return text.encode("utf-8")


def render_md(text: str, name: str) -> bytes:
    return text.encode("utf-8")


def render_rtf(text: str, name: str) -> bytes:
    def esc(value: str) -> str:
        return (
            value.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}").replace("\r", " ")
        )

    body = "\\par\n".join(esc(line) for line in text.splitlines())
    return (
        f"{{\\rtf1\\ansi\\deff0{{\\fonttbl{{\\f0\\fnil Consolas;}}}}\\f0\\fs20 {body}\\par}}\n"
    ).encode()


def render_docx(text: str, name: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    for idx, (line, heading) in enumerate(_to_runs(text)):
        if heading:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line.upper())
            run.bold = True
            run.font.size = Pt(11.5)
        elif _BULLET_RE.match(line):
            document.add_paragraph(line, style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            if idx < 4:
                run.font.size = Pt(12)
                run.font.name = "Georgia"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_pdf(text: str, name: str) -> bytes:
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
    from reportlab.lib.styles import (  # type: ignore[import-untyped]
        ParagraphStyle,
        getSampleStyleSheet,
    )
    from reportlab.lib.units import cm  # type: ignore[import-untyped]
    from reportlab.platypus import (  # type: ignore[import-untyped]
        ListFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.9 * cm,
        rightMargin=1.9 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.6 * cm,
    )
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "CVNormal", parent=styles["BodyText"], fontSize=9.5, leading=12.5, spaceAfter=2
    )
    heading = ParagraphStyle(
        "CVHeading", parent=styles["Heading3"], fontSize=11, leading=14, spaceBefore=8, spaceAfter=3
    )
    lead = ParagraphStyle("CVLead", parent=normal, fontSize=11, leading=14, spaceAfter=4)

    story: list[object] = []
    for idx, (line, is_heading) in enumerate(_to_runs(text)):
        bullet = _BULLET_RE.match(line)
        if is_heading:
            story.append(Paragraph(line.replace("&", "&amp;"), heading))
        elif bullet:
            item = Paragraph(line[bullet.end() :], normal)
            story.append(ListFlowable([item], bulletType="bullet"))
        else:
            story.append(Paragraph(line.replace("&", "&amp;"), lead if idx < 4 else normal))
    story.append(Spacer(1, 6))
    doc.build(story)
    return buffer.getvalue()


Renderer: TypeAlias = Callable[[str, str], bytes]

_EXTENSION_RENDERERS: dict[str, Renderer] = {
    ".txt": render_txt,
    ".md": render_md,
    ".rtf": render_rtf,
    ".docx": render_docx,
    ".pdf": render_pdf,
}
