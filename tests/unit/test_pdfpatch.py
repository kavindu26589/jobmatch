from __future__ import annotations

import io
import re

from reportlab.pdfgen import canvas as rlcanvas

from jobmatcher.pdfpatch import diff_texts, extract_word_boxes, patch_binary
from jobmatcher.resume.parsers import extract_text


def _make_pdf(path, lines) -> None:
    c = rlcanvas.Canvas(str(path), pagesize=(595.0, 842.0))
    y = 800.0
    for line, size in lines:
        c.setFont("Helvetica", size)
        c.drawString(72, y, line)
        y -= 24
    c.showPage()
    c.save()


def test_diff_texts_reports_word_hunks():
    assert diff_texts("same", "same") == []
    changes = diff_texts("Python Developer", "Python Engineer")
    assert [(c.original, c.rewritten) for c in changes] == [("Developer", "Engineer")]


def test_diff_texts_ignores_pure_whitespace_and_insertions():
    assert diff_texts("A B\nC", "A B C") == []
    assert diff_texts("A B", "A B C") == []


def test_diff_texts_merges_adjacent_word_changes():
    before = "Python and SQL"
    after = "Python, SQL and Docker"
    changes = diff_texts(before, after)
    assert changes
    assert all(c.original.strip() for c in changes)
    rebuilt = before
    for change in changes:
        rebuilt = rebuilt.replace(change.original, change.rewritten, 1)
    assert rebuilt == after


def test_patch_pdf_replaces_phrase_in_place(tmp_path):
    src = tmp_path / "cv.pdf"
    _make_pdf(src, [("AI / ML Developer with Python", 12.0), ("Python Engineer", 12.0)])
    original = extract_text(src)
    edited = original.replace("Developer", "Engineer")
    result = patch_binary(src.read_bytes(), original, edited, "pdf")
    assert result.applied == 1 and result.total == 1
    assert not result.notes
    with io.BytesIO(result.data) as stream:
        import pymupdf

        with pymupdf.open(stream=stream, filetype="pdf") as doc:
            assert len(doc) == 1
            text = doc[0].get_text()
            assert "Engineer" in text
            assert "Developer" not in text
            assert "AI / ML" in text


def test_patch_pdf_rewrites_multi_word_phrase_in_place(tmp_path):
    src = tmp_path / "cv.pdf"
    _make_pdf(src, [("AI / ML Developer with Python", 12.0)])
    original = extract_text(src)
    edited = original.replace("Developer with", "Engineer at")
    result = patch_binary(src.read_bytes(), original, edited, "pdf")
    assert result.applied == 1 and not result.notes
    with io.BytesIO(result.data) as stream:
        import pymupdf

        with pymupdf.open(stream=stream, filetype="pdf") as doc:
            text = doc[0].get_text()
            assert "Engineer at" in text
            assert "Python" in text
            assert "Developer with" not in text


def test_patch_pdf_reports_unplaced_changes(tmp_path):
    src = tmp_path / "cv.pdf"
    _make_pdf(src, [("AI / ML Developer with Python", 12.0)])
    result = patch_binary(src.read_bytes(), "Ruby Specialist", "Go Specialist", "pdf")
    assert result.applied == 0 and result.total == 1
    assert any("could not place" in note for note in result.notes)


def test_patch_pdf_repeated_phrase_consumes_distinct_glyphs(tmp_path):
    """Regression: when the same phrase appears several times, each identical
    change must bind to its OWN glyph occurrence. Before the fix, hunks for
    ``"Developer"`` all located the first occurrence (case-insensitive +
    non-consuming), so only one of several identical phrases was ever
    rewritten in place -- the rest stayed untouched."""
    src = tmp_path / "cv.pdf"
    _make_pdf(
        src,
        [
            ("AI / ML Developer with Python", 12.0),
            ("Python Developer for 6 years", 12.0),
            ("Working Developers on GitHub", 12.0),
        ],
    )
    original = extract_text(src)
    edited = re.sub(r"Developer(s?)", r"ENGINEER\1", original)
    assert len(diff_texts(original, edited)) == 3  # every occurrence is a change

    result = patch_binary(src.read_bytes(), original, edited, "pdf")
    assert result.applied == 3 == result.total
    assert not result.notes

    with io.BytesIO(result.data) as stream:
        import pymupdf

        with pymupdf.open(stream=stream, filetype="pdf") as doc:
            text = doc[0].get_text()
    assert text.count("ENGINEER") == 3
    assert "Developer" not in text
    assert "Python" in text and "GitHub" in text


def test_patch_pdf_returns_original_bytes_when_no_changes(tmp_path):
    src = tmp_path / "cv.pdf"
    _make_pdf(src, [("AI / ML Developer with Python", 12.0)])
    data = src.read_bytes()
    original = extract_text(src)
    assert patch_binary(data, original, original, "pdf").data == data
    assert patch_binary(data, original, original + "\nNew appended section", "pdf").data == data


def test_extract_word_boxes_segments_words(tmp_path):
    src = tmp_path / "cv.pdf"
    _make_pdf(src, [("AI / ML Developer with Python", 12.0)])
    pages = extract_word_boxes(src)
    assert len(pages) == 1
    words = [w.text for w in pages[0].words]
    assert "AI" in words and "Developer" in words
    hit = next(w for w in pages[0].words if w.text == "Developer")
    assert hit.x0 < hit.x1 and hit.y0 < hit.y1


def test_patch_docx_replaces_across_runs(tmp_path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    src = tmp_path / "cv.docx"
    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("AI / ML ").bold = False
    paragraph.add_run("Developer").bold = True
    paragraph.add_run(" with 5 years")
    document.save(str(src))

    original = extract_text(src)
    edited = original.replace("Developer", "Engineer")
    result = patch_binary(src.read_bytes(), original, edited, "docx")
    assert result.applied == 1 and not result.notes

    reparsed = docx.Document(io.BytesIO(result.data))
    paragraph = reparsed.paragraphs[0]
    assert paragraph.text == "AI / ML Engineer with 5 years"
    assert paragraph.runs[0].bold is False
    assert paragraph.runs[1].bold is True
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_patch_docx_is_case_insensitive(tmp_path):
    import docx

    src = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Senior Python Developer")
    document.save(str(src))

    results = patch_binary(
        src.read_bytes(), "Senior Python developer", "Senior Python Engineer", "docx"
    )
    assert results.applied == 1
    reparsed = docx.Document(io.BytesIO(results.data))
    assert reparsed.paragraphs[0].text == "Senior Python Engineer"


def test_patch_docx_reports_unplaced_changes(tmp_path):
    import docx

    src = tmp_path / "cv.docx"
    document = docx.Document()
    document.add_paragraph("Python and SQL")
    document.save(str(src))

    results = patch_binary(src.read_bytes(), "Ruby Specialist", "Go Specialist", "docx")
    assert results.applied == 0 and results.total == 1
    assert any("could not place" in note for note in results.notes)
